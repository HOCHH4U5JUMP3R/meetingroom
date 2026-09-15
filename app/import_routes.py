from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook

from . import main

app = main.app
Room = main.Room
YearlyBudget = main.YearlyBudget

# Keep a useful planning/history window available even before the first budget is entered.
def seed_budget_years():
    session = main.SessionLocal()
    try:
        current = date.today().year
        years = range(current - 5, current + 7)
        existing = {(x.year, x.site) for x in session.scalars(select(YearlyBudget)).all()}
        for year in years:
            if (year, None) not in existing:
                session.add(YearlyBudget(year=year, site=None, budget=0))
        session.commit()
    finally:
        session.close()

seed_budget_years()

ROOM_HEADERS = [
    'Raumname', 'Standort/Ort', 'Gebäude', 'Etage', 'Raumnummer',
    'Länge (m)', 'Breite (m)', 'Höhe (m)', 'Sitzplätze', 'Besonderheit',
    'Kategorie', 'Outlook-Ressource', 'Anschlüsse', 'Verantwortlicher',
    'Host-Name', 'Status', 'Letzte Modernisierung', 'Anmerkungen'
]

HEADER_ALIASES = {
    'raumname': 'name', 'name': 'name',
    'standort/ort': 'site', 'standort': 'site', 'ort': 'site',
    'gebäude': 'building', 'gebaeude': 'building',
    'etage': 'floor', 'stockwerk': 'floor',
    'raumnummer': 'room_number', 'raum-nr.': 'room_number', 'raum nr': 'room_number',
    'länge (m)': 'length', 'laenge (m)': 'length', 'länge': 'length', 'laenge': 'length',
    'breite (m)': 'width', 'breite': 'width',
    'höhe (m)': 'height', 'hoehe (m)': 'height', 'höhe': 'height', 'hoehe': 'height',
    'sitzplätze': 'seats', 'sitzplaetze': 'seats', 'sitzplätze ': 'seats',
    'besonderheit': 'specialty', 'kategorie': 'category',
    'outlook-ressource': 'outlook_resource', 'outlookressource': 'outlook_resource',
    'anschlüsse': 'connections', 'anschluesse': 'connections',
    'verantwortlicher': 'owner', 'host-name': 'host_name', 'hostname': 'host_name',
    'status': 'status', 'letzte modernisierung': 'last_modernization',
    'modernisierung': 'last_modernization', 'anmerkungen': 'notes', 'notizen': 'notes'
}


def clean(v):
    if v is None:
        return ''
    return str(v).strip()


def number(v):
    if v in (None, ''):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = clean(v).replace(' ', '').replace(',', '.')
    try:
        return float(text)
    except ValueError:
        return None


def integer(v):
    n = number(v)
    return int(n) if n is not None else None


def parse_date(v):
    if v in (None, ''):
        return None
    if hasattr(v, 'date') and not isinstance(v, datetime):
        return v
    if isinstance(v, datetime):
        return v.date()
    text = clean(v)
    for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def normalize_header(v):
    return clean(v).casefold().replace('  ', ' ')


def row_to_room(row, headers):
    values = {HEADER_ALIASES.get(normalize_header(h)): row[i] for i, h in enumerate(headers) if HEADER_ALIASES.get(normalize_header(h))}
    name = clean(values.get('name'))
    if not name:
        raise ValueError('Raumname fehlt')
    return {
        'name': name,
        'site': clean(values.get('site')),
        'building': clean(values.get('building')),
        'floor': clean(values.get('floor')),
        'room_number': clean(values.get('room_number')),
        'length': number(values.get('length')),
        'width': number(values.get('width')),
        'height': number(values.get('height')),
        'seats': integer(values.get('seats')),
        'specialty': clean(values.get('specialty')),
        'category': clean(values.get('category')),
        'outlook_resource': clean(values.get('outlook_resource')),
        'connections': clean(values.get('connections')),
        'owner': clean(values.get('owner')),
        'host_name': clean(values.get('host_name')),
        'status': clean(values.get('status')) or 'Aktiv',
        'last_modernization': parse_date(values.get('last_modernization')),
        'notes': clean(values.get('notes')),
    }


@app.get('/api/rooms/import-template')
def room_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Räume'
    ws.append(ROOM_HEADERS)
    ws.append([
        'Beispiel Nürnberg 101', 'Nürnberg', 'Gebäude A', '1', '101',
        8.5, 5.2, 2.8, 12, '', 'Meetingraum', '', 'HDMI, USB-C', '', '', 'Aktiv', '', ''
    ])
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = [28, 20, 18, 12, 14, 12, 12, 12, 12, 30, 20, 28, 30, 24, 24, 14, 22, 40]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx) if idx <= 26 else 'A'].width = width
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return StreamingResponse(
        bio,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="meetingroom-raum-import.xlsx"'}
    )


@app.post('/api/rooms/import')
async def import_rooms(file: UploadFile = File(...), db: Session = Depends(main.db)):
    if not file.filename or not file.filename.lower().endswith(('.xlsx', '.xlsm')):
        raise HTTPException(400, 'Bitte eine Excel-Datei (.xlsx oder .xlsm) hochladen.')
    try:
        raw = await file.read()
        wb = load_workbook(BytesIO(raw), data_only=True)
    except Exception as exc:
        raise HTTPException(400, f'Excel-Datei konnte nicht gelesen werden: {exc}')

    ws = wb['Räume'] if 'Räume' in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, 'Die Excel-Datei enthält keine Daten.')
    headers = list(rows[0])
    if not any(HEADER_ALIASES.get(normalize_header(h)) == 'name' for h in headers):
        raise HTTPException(400, 'Spalte „Raumname“ fehlt. Nutze am besten die bereitgestellte Vorlage.')

    created = 0
    updated = 0
    errors = []
    seen = set()
    for excel_row, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, '') for v in row):
            continue
        try:
            values = row_to_room(row, headers)
            if values['name'].casefold() in seen:
                raise ValueError('Raumname kommt in der Datei doppelt vor')
            seen.add(values['name'].casefold())
            existing = db.scalar(select(Room).where(Room.name == values['name']))
            if existing:
                for key, value in values.items():
                    setattr(existing, key, value)
                updated += 1
            else:
                db.add(Room(**values))
                created += 1
        except Exception as exc:
            errors.append({'row': excel_row, 'error': str(exc)})

    if errors and created == 0 and updated == 0:
        db.rollback()
        raise HTTPException(400, {'message': 'Keine Räume importiert.', 'errors': errors})
    db.commit()
    return {
        'created': created,
        'updated': updated,
        'errors': errors,
        'total_processed': created + updated,
        'message': f'{created} neue Räume angelegt, {updated} bestehende Räume aktualisiert.'
    }
